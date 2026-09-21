#!/usr/bin/env python3
"""
Converte a planilha da TBCA no arquivo que a ferramenta importa.

    python3 converter-tbca.py tbca_alimentos_selecionados.xlsx tbca.json

POR QUE O RESULTADO NÃO ENTRA NO REPOSITÓRIO

A TBCA é CC BY-NC-ND 4.0, e a aba "Fonte e notas" da planilha diz "sem fins
comerciais... Não redistribuir". NC é não comercial, ND é sem derivados. O
site é público e o serviço é pago: publicar a base ali seria distribuir uma
cópia. O arquivo gerado é dela, vai para o computador dela, e ela o carrega
na ferramenta — que guarda tudo no navegador. O script fica versionado; a
saída, não.

REGRAS QUE ESTE CONVERSOR SEGUE

  * "tr" (traço), "NA" (não analisado) e "-" (sem informação) viram NULO,
    nunca zero. Traço não é zero, e tratá-lo como zero mudaria a conta de
    sódio e de ferro de dezenas de alimentos sem nada avisando;
  * zero de verdade continua zero — colesterol de vegetal é zero medido;
  * o código da TBCA é preservado, para dar para auditar linha por linha;
  * a citação e a licença viajam junto com os dados, dentro do arquivo.
"""

import json
import sys

import openpyxl

# Coluna da planilha -> nome interno. O que não está aqui não é importado:
# a ferramenta calcula com estes, e carregar 45 colunas para usar 13 só
# engorda o que fica na memória do navegador dela.
COLUNAS = {
    "Energia (kcal)": "energia_kcal",
    "Carboidrato disponível (g)": "carboidrato",
    "Proteína (g)": "proteina",
    "Lipídios (g)": "lipideos",
    "Fibra alimentar (g)": "fibra_alimentar",
    "Colesterol (mg)": "colesterol",
    "Sódio (mg)": "sodio",
    "Cálcio (mg)": "calcio",
    "Ferro (mg)": "ferro",
    "Potássio (mg)": "potassio",
    "Magnésio (mg)": "magnesio",
    "Zinco (mg)": "zinco",
    "Vitamina C (mg)": "vitamina_c",
}

SEM_VALOR = {"tr", "na", "-", "", "nd"}


def numero(v):
    """Número, ou None. Ver a regra do traço no cabeçalho."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = str(v).strip().replace(",", ".")
    if t.lower() in SEM_VALOR:
        return None
    try:
        return float(t)
    except ValueError:
        return None


def medidas_da_celula(texto):
    """'Colher sopa cheia (20 g) | Escumadeira cheia (85 g)' -> duas medidas.

    O peso é sempre o ÚLTIMO parêntese, e não o primeiro: a TBCA escreve
    "Pedaço/Unidade/Fatia (M) (60 g)", onde o primeiro parêntese é o
    tamanho. Lendo o primeiro, essa medida se perdia calada — e a tapioca,
    que tem só ela, ficava sem medida nenhuma.
    """
    saida = []
    if not texto:
        return saida
    for pedaco in str(texto).split("|"):
        pedaco = pedaco.strip()
        if not pedaco or "(" not in pedaco:
            continue
        nome, _, resto = pedaco.rpartition("(")
        gramas = numero(resto.replace("g", "").replace(")", "").strip())
        nome = nome.strip()
        if not nome or gramas is None or gramas <= 0:
            continue
        saida.append({"nome": nome, "gramas": gramas})
    return saida


def converter(caminho_xlsx):
    wb = openpyxl.load_workbook(caminho_xlsx, data_only=True)
    ws = wb["Alimentos"]

    linhas = list(ws.iter_rows(values_only=True))
    cabecalho = [str(c).strip() if c is not None else "" for c in linhas[0]]
    onde = {nome: cabecalho.index(nome) for nome in COLUNAS if nome in cabecalho}
    faltando = [nome for nome in COLUNAS if nome not in onde]
    if faltando:
        # Parar é melhor que importar meia tabela: uma coluna que mudou de
        # nome sai como "sem informação" em todos os alimentos de uma vez.
        raise SystemExit(
            "A planilha não tem estas colunas, que eu esperava:\n  - "
            + "\n  - ".join(faltando)
        )

    caseiras = {}
    if "Medidas caseiras" in wb.sheetnames:
        for linha in list(wb["Medidas caseiras"].iter_rows(values_only=True))[1:]:
            if linha and linha[0]:
                caseiras[str(linha[0]).strip()] = medidas_da_celula(
                    linha[2] if len(linha) > 2 else None
                )

    alimentos = []
    for linha in linhas[1:]:
        codigo = str(linha[0]).strip() if linha[0] else ""
        nome = str(linha[1]).strip() if len(linha) > 1 and linha[1] else ""
        if not codigo or not nome:
            continue  # a última linha da planilha é a legenda, não um alimento
        alimento = {
            "codigo": codigo,
            "nome": nome,
            "grupo": str(linha[2]).strip() if len(linha) > 2 and linha[2] else "",
            "descricao": str(linha[3]).strip() if len(linha) > 3 and linha[3] else "",
            "medidas": caseiras.get(codigo, []),
        }
        for coluna, chave in COLUNAS.items():
            alimento[chave] = numero(linha[onde[coluna]])
        alimentos.append(alimento)

    notas = [
        str(r[0]).strip()
        for r in wb["Fonte e notas"].iter_rows(values_only=True)
        if r and r[0]
    ]
    citacao = next((n for n in notas if n.startswith("Tabela Brasileira")), "")
    licenca = next((n for n in notas if "licen" in n.lower()), "")
    obtido = next((n for n in notas if "consulta em" in n), "")

    return {
        "fonte": {
            "sigla": "TBCA",
            "nome": "Tabela Brasileira de Composição de Alimentos (USP/FoRC)",
            "versao": "7.3",
            "citacao": citacao,
            "licenca": licenca,
            "obtido_em": obtido,
        },
        "alimentos": alimentos,
    }


def main():
    if len(sys.argv) != 3:
        raise SystemExit("uso: converter-tbca.py <planilha.xlsx> <saida.json>")
    tabela = converter(sys.argv[1])
    with open(sys.argv[2], "w", encoding="utf-8") as saida:
        json.dump(tabela, saida, ensure_ascii=False, indent=1)

    com_medida = sum(1 for a in tabela["alimentos"] if a["medidas"])
    sem_kcal = [a["nome"] for a in tabela["alimentos"] if a["energia_kcal"] is None]
    print(f"{len(tabela['alimentos'])} alimentos, {com_medida} com medida caseira.")
    if sem_kcal:
        print("SEM CALORIA (confira antes de usar): " + ", ".join(sem_kcal))


if __name__ == "__main__":
    main()
